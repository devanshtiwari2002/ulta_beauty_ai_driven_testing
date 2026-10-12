from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent
SCENARIO_FILE = ROOT / "test-artifacts" / "Test_Scenarios.xlsx"
TEST_CASE_FILE = ROOT / "test-artifacts" / "Test_Cases_qTest.xlsx"

expected_headers = [
    "Test Case ID", "Test Scenario ID", "Acceptance Criteria", "Test Case Name",
    "Test Type", "Priority", "Preconditions", "Test Data", "Step No", "Test Step",
    "Expected Result", "Automation Status", "Actual Result", "Status", "Defect ID", "Comments",
]

scenario_wb = load_workbook(SCENARIO_FILE, data_only=True)
scenario_ws = scenario_wb["Test Scenarios"]
approved = {
    row[0].value: {
        "ac": row[1].value,
        "title": row[2].value,
        "type": row[4].value,
        "priority": row[5].value,
    }
    for row in scenario_ws.iter_rows(min_row=2)
    if row[0].value and str(row[0].value).startswith("SC-")
}

case_wb = load_workbook(TEST_CASE_FILE, data_only=True)
case_ws = case_wb["Test Cases"]
actual_headers = [case_ws.cell(1, col).value for col in range(1, 17)]
assert actual_headers == expected_headers, actual_headers

rows = list(case_ws.iter_rows(min_row=2, values_only=True))
groups = defaultdict(list)
for row in rows:
    groups[row[0]].append(row)

assert len(rows) == 190, len(rows)
assert len(groups) == 34, len(groups)
assert len({row[0] for row in rows}) == 34
assert len({row[1] for row in rows}) == 34
assert set(row[1] for row in rows) == set(approved)

for test_case_id, group in groups.items():
    scenario_ids = {row[1] for row in group}
    assert len(scenario_ids) == 1, (test_case_id, scenario_ids)
    scenario_id = next(iter(scenario_ids))
    assert scenario_id in approved
    assert {row[2] for row in group} == {approved[scenario_id]["ac"]}
    assert all(row[8] == index for index, row in enumerate(group, 1))
    assert all(row[9] and row[10] for row in group)
    assert all(row[12] is None and row[13] == "Not Executed" and row[14] is None for row in group)
    assert all(row[7] for row in group)
    assert all(row[6] for row in group)
    assert all(row[3] for row in group)

case_records = []
for test_case_id, group in groups.items():
    scenario_id = next(row[1] for row in group)
    case_records.append({
        "test_case_id": test_case_id,
        "scenario_id": scenario_id,
        "ac": approved[scenario_id]["ac"],
        "scenario_type": approved[scenario_id]["type"],
        "automation": next(row[11] for row in group),
    })

ac_counts = Counter(record["ac"] for record in case_records)
automation_counts = Counter(record["automation"] for record in case_records)
intent_counts = Counter(
    "Boundary" if record["scenario_type"] == "Boundary Value"
    else "Negative" if record["scenario_type"] == "Negative"
    else "Positive"
    for record in case_records
)
manual_scenarios = sorted(
    record["scenario_id"]
    for record in case_records
    if record["automation"] == "Manual"
)

result = {
    "total_scenarios": len(approved),
    "total_test_cases": len(groups),
    "step_rows": len(rows),
    "test_cases_by_AC": dict(sorted(ac_counts.items())),
    "automated_vs_manual": dict(automation_counts),
    "positive_negative_boundary": dict(intent_counts),
    "manual_scenarios": manual_scenarios,
    "verified": True,
}
print("FINAL TRACEABILITY VERIFICATION PASSED")
print(result)
