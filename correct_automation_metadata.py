from pathlib import Path

from openpyxl import load_workbook

path = Path("test-artifacts/Test_Cases_qTest.xlsx")
wb = load_workbook(path, data_only=False)
ws = wb["Test Cases"]

approach_by_ac = {
    "AC1": "UI Automation: Selenium + Cucumber + Page Object Model.",
    "AC2": "UI Automation: Selenium + Cucumber + Page Object Model.",
    "AC3": "UI Automation: Selenium + Cucumber + Page Object Model.",
    "AC4": "UI Automation: Selenium + Cucumber + Page Object Model.",
    "AC5": "Business Rule / Boundary Automation: Java automated validation using controlled recommendation datasets.",
    "AC6": "Business Rule / Boundary Automation: Java automated validation using controlled recommendation datasets.",
    "AC7": "Business Rule / Boundary Automation: Java automated validation using controlled recommendation datasets.",
}

for row in range(2, ws.max_row + 1):
    scenario_id = ws.cell(row, 2).value
    ac = ws.cell(row, 3).value
    if ac not in approach_by_ac:
        raise ValueError(f"Unsupported acceptance criteria in {scenario_id}: {ac}")

    ws.cell(row, 12).value = "Automated"
    ws.cell(row, 13).value = None
    ws.cell(row, 14).value = "Not Executed"
    ws.cell(row, 15).value = None
    ws.cell(row, 16).value = (
        f"Automation approach: {approach_by_ac[ac]} "
        f"Scenario: {scenario_id}; execution evidence pending."
    )

wb.save(path)
print(f"Updated automation metadata in {path}")
