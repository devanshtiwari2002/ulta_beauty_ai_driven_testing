from openpyxl import load_workbook
from pathlib import Path

path = Path("test-artifacts/Test_Cases_qTest.updated.xlsx")
wb = load_workbook(path, data_only=False)
ws = wb["Test Cases"]
exact = {
    "SC-AC5-001": (
        "Fail",
        "Fresh report failure: parameterized product recommendation scenario for vitamin C serum timed out waiting for visibility of //h2[normalize-space()='Similar items for you'] after 20 seconds.",
        "Evidence: target/surefire-reports/testng-results.xml; no defect ID was recorded.",
    ),
    "SC-AC6-001": (
        "Fail",
        "Fresh report failure: cart recommendation relevance was 0.0% versus required 25.0%; 12 recommendations were evaluated and 0 were relevant.",
        "Evidence: target/surefire-reports/testng-results.xml; no defect ID was recorded.",
    ),
    "SC-AC7-001": (
        "Pass",
        "Fresh report passed: 36 recommendations were collected, 23 were in the beauty/skincare/makeup category, and relevance was 63.9% versus required 20.0%.",
        "Evidence: target/surefire-reports/testng-results.xml; no defect ID was recorded.",
    ),
}

for row in range(2, ws.max_row + 1):
    scenario_id = ws.cell(row, 2).value
    if scenario_id in exact:
        status, actual_result, comments = exact[scenario_id]
        ws.cell(row, 13).value = actual_result
        ws.cell(row, 14).value = status
        ws.cell(row, 15).value = None
        ws.cell(row, 16).value = comments
    else:
        ws.cell(row, 13).value = None
        ws.cell(row, 14).value = "Not Executed"
        ws.cell(row, 15).value = None
        ws.cell(row, 16).value = (
            "No direct execution evidence in the fresh report; this row is not correlated "
            "to a recorded Cucumber/TestNG scenario."
        )

wb.save(path)
print(f"narrowed {path}")
